"""Static checks of a generated workbook against its excel-input.json.

The workbook holds formula strings without cached values, so the checks compare
each formula cell with the pattern documented in references/excel-schema.md, at
the rows the layout of the JSON input puts it. They catch hardcoded values in
formula cells, broken references, wrong rollup ranges, empty input cells, a
stale workbook (JSON changed after generation) and a wrong sheet order.

``recalc_check`` optionally recalculates a copy with LibreOffice, when it is on
PATH, and compares the computed values with the figures from helpers.figures.
"""
from __future__ import annotations

import contextlib
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from helpers.config_compat import normalize_config
from helpers.figures import compute_summary, is_number
from helpers.i18n import t
from helpers.wb_summary import _quote_sheet

NUM = r"(-?\d+(?:\.\d+)?(?:e-?\d+)?)"
ERROR_TOKENS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
INPUT_COLS = ("E", "F", "G", "I", "J", "K")
SUMMARY_FROM_WBS = dict(zip("ABCDEFGHIJK", ["B", None, "E", "F", "G", "H", "I", "J", "K", "L", "M"]))
PRIORITY = '=IF(G{r}>=15,"CRITICAL",IF(G{r}>=10,"HIGH",IF(G{r}>=5,"MEDIUM","LOW")))'


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def check_workbook(path: str | Path, data: dict, summary: dict | None = None) -> list[str]:
    """Return a list of errors ("Sheet!Cell: expected ..., found ..."); empty means pass."""
    summary = summary or compute_summary(data)
    with contextlib.redirect_stderr(io.StringIO()):  # the legacy warning is shown once, by compute_summary
        data = normalize_config(data)
    cfg = data["config"]
    lang = cfg.get("lang", "en")
    wb = load_workbook(path, data_only=False)
    names = [t(lang, k) for k in ("sheet_wbs", "sheet_resource_plan", "sheet_risks", "sheet_summary")]
    if wb.sheetnames != names:
        return [f"sheets: expected {names} in this order, found {wb.sheetnames}"]
    errors = _scan_error_tokens(wb)
    layout = wbs_layout(data["phases"])
    _check_wbs(wb[names[0]], layout, errors)
    _check_resource_plan(wb[names[1]], summary, errors)
    risk_total = _check_risks(wb[names[2]], data, layout["total_row"], errors)
    _check_summary(wb[names[3]], data, summary, layout, (names[2], risk_total), errors)
    return errors


def wbs_layout(phases: list[dict]) -> dict[str, Any]:
    """Row numbers of each phase, work package and leaf, as the WBS builder assigns them."""
    row, out = 2, []
    for phase in phases:
        ph = {"row": row, "wps": [], "leaf_rows": []}
        row += 1
        for wp in phase.get("work_packages", []):
            first = row + 1
            row = first + len(wp.get("activities", []))
            leaf_rows = list(range(first, row))
            ph["wps"].append({"row": first - 1, "first": first, "last": row - 1, "leaf_rows": leaf_rows})
            ph["leaf_rows"].extend(leaf_rows)
        out.append(ph)
    return {"phases": out, "total_row": row}


def recalc_check(path: str | Path, data: dict, summary: dict, timeout: int = 120) -> dict[str, Any]:
    """Recalculate a copy with LibreOffice and compare key values with the computed figures."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"status": "skipped", "reason": "LibreOffice (soffice) not found on PATH"}
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "profile").as_uri()
        cmd = [soffice, f"-env:UserInstallation={profile}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", str(Path(tmp, "out")), str(path)]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return {"status": "error", "reason": f"LibreOffice timed out after {timeout} s"}
        converted = Path(tmp, "out", Path(path).stem + ".xlsx")
        if proc.returncode != 0 or not converted.exists():
            return {"status": "error", "reason": (proc.stderr or proc.stdout).strip()[-500:]}
        mismatches = _compare_values(load_workbook(converted, data_only=True), data, summary)
    return {"status": "failed" if mismatches else "passed", "mismatches": mismatches}


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------

def _expect(ws, ref: str, expected: Any, errors: list[str]) -> None:
    actual = ws[ref].value
    if actual != expected:
        errors.append(f"{ws.title}!{ref}: expected {expected!r}, found {actual!r}")


def _expect_match(ws, ref: str, pattern: str, numbers: tuple, errors: list[str]) -> None:
    """Formula must match *pattern*; its captured numbers must equal *numbers* (stale check)."""
    actual = ws[ref].value
    m = re.fullmatch(pattern, actual) if isinstance(actual, str) else None
    if not m:
        errors.append(f"{ws.title}!{ref}: expected a formula like {pattern!r}, found {actual!r}")
        return
    found = tuple(float(g) for g in m.groups())
    if any(abs(a - float(b)) > 1e-9 for a, b in zip(found, numbers)):
        errors.append(f"{ws.title}!{ref}: formula uses {found}, the input says {numbers}")


def _expect_number(ws, ref: str, errors: list[str]) -> None:
    v = ws[ref].value
    if not is_number(v):
        errors.append(f"{ws.title}!{ref}: expected a number (input cell), found {v!r}")


def _scan_error_tokens(wb) -> list[str]:
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and any(tok in v for tok in ERROR_TOKENS):
                    errors.append(f"{ws.title}!{cell.coordinate}: contains an error token: {v!r}")
    return errors


def _sum_list(col: str, rows: list[int]) -> str:
    return "=SUM(" + ",".join(f"{col}{r}" for r in rows) + ")"


# ---------------------------------------------------------------------------
# Sheet checks
# ---------------------------------------------------------------------------

def _pert_row(ws, r: int, errors: list[str]) -> None:
    _expect(ws, f"H{r}", f"=(E{r}+4*F{r}+G{r})/6", errors)
    _expect(ws, f"L{r}", f"=(I{r}+4*J{r}+K{r})/6", errors)
    _expect(ws, f"M{r}", f"=(K{r}-I{r})/6", errors)


def _check_wbs(ws, layout: dict, errors: list[str]) -> None:
    for ph in layout["phases"]:
        for wp in ph["wps"]:
            for r in wp["leaf_rows"]:
                for c in INPUT_COLS:
                    _expect_number(ws, f"{c}{r}", errors)
                _pert_row(ws, r, errors)
                _expect(ws, f"S{r}", f'=IF(R{r}="Y",H{r},0)', errors)
            r, first, last = wp["row"], wp["first"], wp["last"]
            for c in INPUT_COLS + ("S",):
                _expect(ws, f"{c}{r}", f"=SUM({c}{first}:{c}{last})", errors)
            _pert_row(ws, r, errors)
        r = ph["row"]
        for c in INPUT_COLS + ("S",):
            _expect(ws, f"{c}{r}", _sum_list(c, ph["leaf_rows"]), errors)
        _pert_row(ws, r, errors)
    tot = layout["total_row"]
    _expect(ws, f"A{tot}", "TOTAL", errors)
    phase_rows = [ph["row"] for ph in layout["phases"]]
    for c in ("E", "F", "G", "H", "I", "J", "K", "L", "M", "S"):
        _expect(ws, f"{c}{tot}", _sum_list(c, phase_rows), errors)


def _check_resource_plan(ws, summary: dict, errors: list[str]) -> None:
    rp = summary["resource_plan"]
    weeks, first_col = rp["total_weeks"], 4
    last = get_column_letter(first_col + weeks - 1)
    total_col = get_column_letter(first_col + weeks)
    grand = 0.0
    for i, code in enumerate(rp["role_codes"]):
        r = 3 + i
        _expect(ws, f"B{r}", code, errors)
        for w in range(weeks):
            ref = f"{get_column_letter(first_col + w)}{r}"
            _expect_number(ws, ref, errors)
            grand += ws[ref].value if is_number(ws[ref].value) else 0
        _expect(ws, f"{total_col}{r}", f"=SUM(D{r}:{last}{r})", errors)
    expected = summary["effort"]["tech_pert"] - rp["skipped_pd"]
    if abs(grand - expected) > 1.0:
        errors.append(f"{ws.title}: role-week cells sum to {grand:.2f} PD, expected {expected:.2f} "
                      "(Tech PERT less activities without resources) within 1 PD")


def _check_risks(ws, data: dict, wbs_total: int, errors: list[str]) -> int:
    cfg, n = data["config"], len(data["risks"])
    for r in range(2, 2 + n):
        _expect_number(ws, f"E{r}", errors)
        _expect_number(ws, f"F{r}", errors)
        _expect(ws, f"G{r}", f"=E{r}*F{r}", errors)
        _expect(ws, f"H{r}", PRIORITY.format(r=r), errors)
    total, reserve = n + 3, n + 4
    _expect(ws, f"L{total}", f"=SUM(L2:L{n + 1})", errors)
    pattern = rf"=\(WBS!H{wbs_total}\*\(1\+{NUM}\+{NUM}\)\+L{total}\)\*{NUM}"
    numbers = (float(cfg.get("pm_overhead_pct") or 0), float(cfg.get("devops_overhead_pct") or 0),
               float(cfg.get("management_reserve_pct", 0.10)))
    _expect_match(ws, f"L{reserve}", pattern, numbers, errors)
    return total


def _check_summary(ws, data: dict, summary: dict, layout: dict, risks: tuple, errors: list[str]) -> None:
    cfg, lang = data["config"], data["config"].get("lang", "en")
    n = len(layout["phases"])
    for i, ph in enumerate(layout["phases"]):
        r = 2 + i
        for col, wbs_col in SUMMARY_FROM_WBS.items():
            if wbs_col:
                _expect(ws, f"{col}{r}", f"=WBS!{wbs_col}{ph['row']}", errors)
    tot = n + 2
    for col in "CDEFGHIJK":
        _expect(ws, f"{col}{tot}", f"=SUM({col}2:{col}{n + 1})", errors)
    rows = dict(zip(["tech", "pm", "devops", "sub", "cont", "low", "mr", "med", "high", "bill", "ratio"],
                    range(tot + 2, tot + 13)))
    _check_effort_block(ws, rows, tot, cfg, lang, risks, layout["total_row"], errors)
    cal = rows["ratio"] + 2
    _expect(ws, f"A{cal}", t(lang, "summary_calendar_duration"), errors)
    _expect(ws, f"B{cal}", summary["calendar_weeks"], errors)


def _check_effort_block(ws, rows: dict, tot: int, cfg: dict, lang: str, risks: tuple,
                        wbs_total: int, errors: list[str]) -> None:
    b = {k: f"B{v}" for k, v in rows.items()}
    for key, label in (("tech", "summary_tech_pert"), ("sub", "summary_subtotal"),
                       ("low", "fascia_bassa"), ("med", "fascia_media"), ("high", "fascia_alta")):
        _expect(ws, f"A{rows[key]}", t(lang, label), errors)
    _expect(ws, b["tech"], f"=F{tot}", errors)
    _expect_match(ws, b["pm"], rf"=B{rows['tech']}\*{NUM}", (float(cfg.get("pm_overhead_pct") or 0),), errors)
    _expect_match(ws, b["devops"], rf"=B{rows['tech']}\*{NUM}",
                  (float(cfg.get("devops_overhead_pct") or 0),), errors)
    _expect(ws, b["sub"], f"=B{rows['tech']}+B{rows['pm']}+B{rows['devops']}", errors)
    _expect(ws, b["cont"], f"={_quote_sheet(risks[0])}!L{risks[1]}", errors)
    _expect(ws, b["low"], f"=B{rows['sub']}+B{rows['cont']}", errors)
    _expect_match(ws, b["mr"], rf"=B{rows['low']}\*{NUM}",
                  (float(cfg.get("management_reserve_pct") or 0),), errors)
    _expect(ws, b["med"], f"=B{rows['low']}+B{rows['mr']}", errors)
    _expect_match(ws, b["high"], rf"=B{rows['med']}\*\(1\+{NUM}\)",
                  (float(cfg.get("alta_uplift_pct") or 0),), errors)
    _expect(ws, b["bill"], f"=WBS!S{wbs_total}", errors)
    _expect(ws, b["ratio"], f"=B{rows['bill']}/B{rows['tech']}", errors)


# ---------------------------------------------------------------------------
# Recalculated values
# ---------------------------------------------------------------------------

def _compare_values(wb, data: dict, summary: dict) -> list[str]:
    lang = data["config"].get("lang", "en")
    layout = wbs_layout(data["phases"])
    n_phases, n_risks = len(layout["phases"]), len(data["risks"])
    tot = n_phases + 2
    e = summary["effort"]
    wanted = [
        (t(lang, "sheet_wbs"), f"H{layout['total_row']}", e["tech_pert"]),
        (t(lang, "sheet_risks"), f"L{n_risks + 4}", e["management_reserve_risks_sheet"]),
    ]
    for key, offset in (("low_band", 7), ("management_reserve", 8), ("medium_band", 9), ("high_band", 10)):
        wanted.append((t(lang, "sheet_summary"), f"B{tot + offset}", e[key]))
    out = []
    for sheet, ref, expected in wanted:
        v = wb[sheet][ref].value
        if not is_number(v) or abs(v - expected) > 0.01:
            out.append(f"{sheet}!{ref}: recalculated {v!r}, computed {expected}")
    return out
