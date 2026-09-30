"""Every bundled example input must still generate a valid workbook.

This catches schema drift: when the generator or the documented schema changes,
an example that no longer fits fails here instead of in a user's session.
"""
import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

SCRIPTS_DIR = Path(__file__).resolve().parent.parent
EXAMPLES_DIR = SCRIPTS_DIR.parent / "examples"
EXAMPLE_INPUTS = sorted(EXAMPLES_DIR.glob("*/excel-input.json")) + [EXAMPLES_DIR / "sample-input.json"]

EXPECTED_SHEETS = {
    "en": ["WBS", "Resource Plan", "Risks", "Summary"],
    "it": ["WBS", "Pianificazione Risorse", "Rischi", "Riepilogo"],
}


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS_DIR / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_examples_are_found():
    assert len(EXAMPLE_INPUTS) >= 7


@pytest.mark.parametrize("src", EXAMPLE_INPUTS, ids=lambda p: p.parent.name if p.name != "sample-input.json" else "sample")
def test_example_generates_and_passes_checks(src, tmp_path):
    gen = _load("generate_excel")
    summ = _load("summarize")
    from helpers import workbook_checks

    data = json.loads(src.read_text(encoding="utf-8"))
    errors, _ = summ.validate_input(data)
    assert errors == [], errors

    out = tmp_path / "out.xlsx"
    assert gen.main(str(src), str(out)) == 0

    wb = load_workbook(out)
    assert wb.sheetnames == EXPECTED_SHEETS[data["config"].get("lang", "en")]
    assert workbook_checks.check_workbook(out, data) == []


def test_sample_respects_8_80_rule():
    summ = _load("summarize")
    data = json.loads((EXAMPLES_DIR / "sample-input.json").read_text(encoding="utf-8"))
    _, warnings = summ.validate_input(data)
    assert not [w for w in warnings if "8/80" in w], warnings


def test_sample_uses_only_keys_the_generator_reads():
    summ = _load("summarize")
    data = json.loads((EXAMPLES_DIR / "sample-input.json").read_text(encoding="utf-8"))
    _, warnings = summ.validate_input(data)
    assert not [w for w in warnings if "not read" in w], warnings
